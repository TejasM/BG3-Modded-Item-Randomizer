"""md2docx.py - render MIR's Markdown docs (README.md / INSTALL.md) to .docx with python-docx.

Replaces the pandoc step used for the CCNS bundle (pandoc is not installed here). Handles the
subset of Markdown the docs use: ATX headings, paragraphs, bullet lists (with nested indent),
numbered lists, fenced code blocks, pipe tables, and inline **bold**, *italic*, `code`,
<https://...> and [text](url) links. Escaped pipes (\\|) inside table cells are unescaped.

Usage:  python md2docx.py <in.md> <out.docx> "<Document title>"
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|<https?://[^>\s]+>|\[[^\]]+?\]\(https?://[^)\s]+\))")

def add_runs(par, text, base_bold=False):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); r.bold = base_bold or None
        tok = m.group(0)
        if tok.startswith("**"):
            add_runs(par, tok[2:-2], base_bold=True)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            if base_bold: r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); r.italic = True
            if base_bold: r.bold = True
        elif tok.startswith("<"):
            add_link(par, tok[1:-1], tok[1:-1])
        else:
            mm = re.match(r"\[([^\]]+?)\]\((https?://[^)\s]+)\)", tok)
            add_link(par, mm.group(1), mm.group(2))
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); r.bold = base_bold or None

def add_link(par, text, url):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    new_run.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); new_run.append(t)
    h.append(new_run); par._p.append(h)

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill); tcPr.append(shd)

def render(md_path, out_path, title):
    lines = open(md_path, encoding="utf-8").read().replace("\r\n", "\n").split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Inches(1))
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    doc.core_properties.title = title; doc.core_properties.author = "Serpentine (NexusMods: SerpentineShel)"
    i = 0; n = len(lines); first_h1 = True
    while i < n:
        ln = lines[i]
        if ln.startswith("```"):
            i += 1; buf = []
            while i < n and not lines[i].startswith("```"): buf.append(lines[i]); i += 1
            i += 1
            for cl in buf:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Inches(0.3)
                r = p.add_run(cl if cl else " "); r.font.name = "Consolas"; r.font.size = Pt(9.5)
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            level = len(m.group(1)); text = m.group(2)
            if level == 1 and first_h1:
                doc.add_heading(text, level=0); first_h1 = False
            else:
                doc.add_heading(text, level=level)
            i += 1; continue
        if ln.startswith("|") and i + 1 < n and re.match(r"^\|\s*-", lines[i + 1]):
            rows = []
            while i < n and lines[i].startswith("|"):
                if not re.match(r"^\|\s*-", lines[i]):
                    cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[i])[1:-1]]
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.cell(ri, ci); cell.text = ""
                    par = cell.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
                    add_runs(par, row[ci] if ci < len(row) else "", base_bold=(ri == 0))
                    for r in par.runs: r.font.size = Pt(9.5)
                    if ri == 0: shade(cell, "E7E6E6")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        mb = re.match(r"^(\s*)[-*]\s+(.*)$", ln)
        if mb:
            indent = len(mb.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            add_runs(p, mb.group(2)); i += 1
            while i < n and lines[i].startswith("  ") and not re.match(r"^\s*[-*]\s+", lines[i]) and lines[i].strip():
                add_runs(p, " " + lines[i].strip()); i += 1
            continue
        mn = re.match(r"^(\d+)\.\s+(.*)$", ln)
        if mn:
            p = doc.add_paragraph(style="List Number"); add_runs(p, mn.group(2)); i += 1
            while i < n and lines[i].startswith("   ") and lines[i].strip() and not lines[i].strip().startswith("```"):
                add_runs(p, " " + lines[i].strip()); i += 1
            continue
        if ln.strip() == "" or ln.strip() == "---":
            i += 1; continue
        # paragraph: join following non-blank, non-structural lines
        buf = [ln.strip()]; i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,3}\s|```|\||\s*[-*]\s|\d+\.\s)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); add_runs(p, " ".join(buf))
    doc.save(out_path)
    print("wrote", out_path)

if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2], sys.argv[3])
